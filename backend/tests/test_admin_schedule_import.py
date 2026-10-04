import io
from datetime import date, timedelta

import pytest

from app.services import schedule_import


TODAY = date.today()
TOMORROW = TODAY + timedelta(days=1)
DAY_AFTER = TODAY + timedelta(days=2)
PAST = TODAY - timedelta(days=3)


def _xlsx(header: list, rows: list[list]) -> bytes:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(header)
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _xlsx_sheets(sheets: list[tuple[str, list, list[list]]]) -> bytes:
    from openpyxl import Workbook

    workbook = Workbook()
    workbook.remove(workbook.active)
    for title, header, rows in sheets:
        sheet = workbook.create_sheet(title)
        if header:
            sheet.append(header)
        for row in rows:
            sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _upload(data: bytes, filename: str = "schedule.xlsx"):
    return {"file": (filename, data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}


# ---- 解析层（不经过 HTTP） ----


def test_parse_basic_columns():
    data = _xlsx(
        ["日期", "开始时间", "结束时间", "标题", "备注", "分类", "优先级", "状态"],
        [
            [TOMORROW.isoformat(), "09:00", "10:00", "晨会", "每周例会", "工作", "高", "待办"],
            [DAY_AFTER.isoformat(), "14:00:00", "15:30", "项目复盘", "", "", "", ""],
        ],
    )
    result = schedule_import.parse_schedule_excel(data)
    assert result["skipped"] == 0
    assert len(result["rows"]) == 2
    first = result["rows"][0]
    assert first["title"] == "晨会"
    assert first["start_time"] == "09:00"
    assert first["end_time"] == "10:00"
    assert first["category"] == "工作"
    assert first["priority"] == "high"
    assert first["status"] == "pending"
    assert result["rows"][1]["category"] == "日程"
    assert result["columns"]["date"] == "日期"


def test_parse_date_and_time_formats():
    data = _xlsx(
        ["日期", "时间", "事项"],
        [
            ["2026年10月5日", "08:00-09:30", "早读"],
            ["2026/10/6", "下午2点半", "社团活动"],
            ["20261007", "20:00", "晚自习"],
        ],
    )
    result = schedule_import.parse_schedule_excel(data)
    assert [row["title"] for row in result["rows"]] == ["早读", "社团活动", "晚自习"]
    assert result["rows"][0]["date"] == "2026-10-05"
    assert result["rows"][0]["start_time"] == "08:00"
    assert result["rows"][0]["end_time"] == "09:30"
    assert result["rows"][1]["date"] == "2026-10-06"
    assert result["rows"][1]["start_time"] == "14:30"
    assert result["rows"][2]["date"] == "2026-10-07"


def test_parse_skips_bad_rows():
    data = _xlsx(
        ["日期", "标题"],
        [
            [TOMORROW.isoformat(), "正常事项"],
            ["不是日期", "坏日期"],
            [TOMORROW.isoformat(), ""],
        ],
    )
    result = schedule_import.parse_schedule_excel(data)
    assert result["skipped"] == 2
    assert len(result["rows"]) == 1
    assert any("坏日期" in warning for warning in result["warnings"])


def test_parse_repeat_expands_dates():
    until = TODAY + timedelta(days=6)
    data = _xlsx(
        ["日期", "标题", "重复", "重复至"],
        [
            [TODAY.isoformat(), "每日晨读", "每天", until.isoformat()],
            [TOMORROW.isoformat(), "每周例会", "每周", (TOMORROW + timedelta(days=14)).isoformat()],
        ],
    )
    result = schedule_import.parse_schedule_excel(data)
    daily = [row for row in result["rows"] if row["title"] == "每日晨读"]
    weekly = [row for row in result["rows"] if row["title"] == "每周例会"]
    assert len(daily) == 7
    assert len(weekly) == 3
    assert weekly[0]["date"] == TOMORROW.isoformat()
    assert all(row["start_time"] is None for row in result["rows"])


def test_parse_priority_and_status_words():
    data = _xlsx(
        ["日期", "标题", "优先级", "状态"],
        [
            [TOMORROW.isoformat(), "答辩准备", "紧急", "进行中"],
            [TOMORROW.isoformat(), "交材料", "低", "已完成"],
            [TOMORROW.isoformat(), "取消的事", "普通", "已取消"],
        ],
    )
    rows = schedule_import.parse_schedule_excel(data)["rows"]
    assert [row["priority"] for row in rows] == ["high", "low", "medium"]
    assert [row["status"] for row in rows] == ["in_progress", "done", "cancelled"]


def test_auto_detect_with_unusual_headers():
    """表头不是「日期/标题」时，按内容自动判断。"""
    data = _xlsx(
        ["上课时间", "课程名称", "上课地点"],
        [
            [TOMORROW.isoformat(), "高等数学", "B202"],
            [DAY_AFTER.isoformat(), "线性代数", "训1223"],
        ],
    )
    result = schedule_import.parse_schedule_excel(data)
    assert [row["title"] for row in result["rows"]] == ["高等数学", "线性代数"]
    assert result["columns"]["date"] == "上课时间"
    assert result["columns"]["title"] == "课程名称"
    assert any("已自动识别" in warning for warning in result["warnings"])


def test_auto_detect_without_header_row():
    """整张表没有表头，直接是数据。"""
    data = _xlsx_sheets([("Sheet1", [], [[TOMORROW.isoformat(), "晨会"], [DAY_AFTER.isoformat(), "整理资料"]])])
    result = schedule_import.parse_schedule_excel(data)
    assert [row["title"] for row in result["rows"]] == ["晨会", "整理资料"]
    assert result["columns"]["date"] == "第 1 列"
    assert result["columns"]["title"] == "第 2 列"


def test_auto_detect_header_like_first_row():
    """第一行像表头但不是已知写法：不应被当成一条日程。"""
    data = _xlsx(
        ["日期安排", "事项内容", "地点"],
        [[TOMORROW.isoformat(), "小组讨论", "A101"]],
    )
    result = schedule_import.parse_schedule_excel(data)
    assert len(result["rows"]) == 1
    assert result["rows"][0]["title"] == "小组讨论"


def test_parse_month_day_without_year():
    data = _xlsx(["日期", "标题"], [["10月5日", "没写年份的日程"]])
    result = schedule_import.parse_schedule_excel(data)
    assert result["rows"][0]["date"].endswith("-10-05")
    assert any("没写年份" in warning for warning in result["warnings"])


def test_parse_skips_non_schedule_first_sheet():
    """第一个工作表是说明页时，自动看下一个工作表。"""
    data = _xlsx_sheets(
        [
            ("说明", ["本表说明"], [["请填写下一个工作表"]]),
            ("日程表", ["日期", "标题"], [[TOMORROW.isoformat(), "晨会"]]),
        ]
    )
    result = schedule_import.parse_schedule_excel(data)
    assert [row["title"] for row in result["rows"]] == ["晨会"]


def test_parse_keeps_warnings_when_nothing_parsed():
    """一行都没解析出来时，仍然返回具体原因，方便管理员改表。"""
    data = _xlsx(["日期", "标题"], [["下周三", "开班会"]])
    result = schedule_import.parse_schedule_excel(data)
    assert result["rows"] == []
    assert result["skipped"] == 1
    assert any("开班会" in warning for warning in result["warnings"])


def test_parse_week_grid_layout():
    """横向按周排布：第一行是日期，下面格子里写当天安排。"""
    from openpyxl import Workbook

    week1 = [TOMORROW + timedelta(days=i) for i in range(7)]
    week2 = [day + timedelta(days=7) for day in week1]
    wb = Workbook()
    ws = wb.active
    ws.append(week1)
    ws.append(["词汇课"] * 7)
    ws.append(["阅读课1", "阅读课2", "", "", "阅读课5", "", ""])
    ws.append([None, None, None, None, "写作课1", None, "写作课2"])
    ws.append([None, None, None, "复习", None, None, None])
    ws.append(week2)
    ws.append(["词汇课"] * 7)
    ws.append([None, "阅读课13", None, None, None, "阅读课14", None])
    buf = io.BytesIO()
    wb.save(buf)

    result = schedule_import.parse_schedule_excel(buf.getvalue())
    titles = [row["title"] for row in result["rows"]]
    assert titles.count("词汇课") == 14
    assert result["skipped"] == 0
    assert any("横向按周排布" in warning for warning in result["warnings"])

    by_date = {}
    for row in result["rows"]:
        by_date.setdefault(row["date"], []).append(row["title"])
    assert by_date[TOMORROW.isoformat()] == ["词汇课", "阅读课1"]
    assert by_date[(TOMORROW + timedelta(days=1)).isoformat()] == ["词汇课", "阅读课2"]
    assert by_date[(TOMORROW + timedelta(days=3)).isoformat()] == ["词汇课", "复习"]
    assert by_date[(TOMORROW + timedelta(days=4)).isoformat()] == ["词汇课", "阅读课5", "写作课1"]
    assert by_date[(TOMORROW + timedelta(days=7)).isoformat()] == ["词汇课"]
    assert by_date[(TOMORROW + timedelta(days=8)).isoformat()] == ["词汇课", "阅读课13"]


def test_grid_cell_with_time_and_multiline():
    from openpyxl import Workbook

    week = [TOMORROW + timedelta(days=i) for i in range(7)]
    wb = Workbook()
    ws = wb.active
    ws.append(week)
    ws.append(["08:00 背单词\n20:00-21:00 真题", "", "", "", "", "", ""])
    buf = io.BytesIO()
    wb.save(buf)

    rows = schedule_import.parse_schedule_excel(buf.getvalue())["rows"]
    assert [(row["start_time"], row["title"]) for row in rows] == [
        ("08:00", "背单词"),
        ("20:00", "真题"),
    ]
    assert rows[1]["end_time"] == "21:00"
    assert {row["date"] for row in rows} == {TOMORROW.isoformat()}


def test_parse_requires_date_and_title_columns():
    with pytest.raises(schedule_import.ScheduleImportError):
        schedule_import.parse_schedule_excel(_xlsx(["姓名", "分数"], [["张三", 90]]))
    with pytest.raises(schedule_import.ScheduleImportError):
        schedule_import.parse_schedule_excel(b"not an excel file")


def test_template_workbook_is_readable():
    data = schedule_import.build_template_workbook()
    result = schedule_import.parse_schedule_excel(data)
    assert len(result["rows"]) >= 1


# ---- HTTP 接口 ----


def test_preview_requires_admin(client, auth_headers):
    data = _xlsx(["日期", "标题"], [[TOMORROW.isoformat(), "晨会"]])
    response = client.post("/api/admin/schedule/preview", files=_upload(data), headers=auth_headers)
    assert response.status_code == 403


def test_preview_and_template(client, admin_headers):
    data = _xlsx(["日期", "标题", "开始时间"], [[TOMORROW.isoformat(), "晨会", "09:00"]])
    response = client.post("/api/admin/schedule/preview", files=_upload(data), headers=admin_headers)
    assert response.status_code == 200
    body = response.json()
    assert body["rows"][0]["title"] == "晨会"
    assert body["columns"]["title"] == "标题"

    template = client.get("/api/admin/schedule/template", headers=admin_headers)
    assert template.status_code == 200
    assert template.content[:2] == b"PK"

    wrong = client.post("/api/admin/schedule/preview", files=_upload(b"hello", "schedule.txt"), headers=admin_headers)
    assert wrong.status_code == 400


def test_import_to_self(client, admin_headers):
    rows = [
        {"date": TOMORROW.isoformat(), "title": "晨会", "start_time": "09:00", "end_time": "10:00", "category": "工作"},
        {"date": PAST.isoformat(), "title": "已过期的日程"},
    ]
    response = client.post(
        "/api/admin/schedule/import",
        json={"rows": rows, "include_self": True},
        headers=admin_headers,
    )
    assert response.status_code == 200
    body = response.json()
    assert body["created"] == 1
    assert body["skipped_past"] == 1
    assert len(body["plan_ids"]) == 1

    plans = client.get("/api/plans", headers=admin_headers).json()
    assert [(plan["date"], plan["title"]) for plan in plans] == [(TOMORROW.isoformat(), "晨会")]

    # 再导一次：同样的日程会被去重
    again = client.post(
        "/api/admin/schedule/import",
        json={"rows": rows, "include_self": True},
        headers=admin_headers,
    ).json()
    assert again["created"] == 0
    assert again["skipped_duplicate"] == 1


def test_import_requires_target(client, admin_headers):
    rows = [{"date": TOMORROW.isoformat(), "title": "晨会"}]
    response = client.post("/api/admin/schedule/import", json={"rows": rows}, headers=admin_headers)
    assert response.status_code == 400

    unknown = client.post(
        "/api/admin/schedule/import",
        json={"rows": rows, "user_ids": [99999]},
        headers=admin_headers,
    )
    assert unknown.status_code == 400


def test_import_to_all_users_and_rollback(client, admin_headers, register_user):
    other_headers = register_user("member@example.com")
    rows = [{"date": TOMORROW.isoformat(), "title": "全员通知", "category": "通知"}]
    response = client.post(
        "/api/admin/schedule/import",
        json={"rows": rows, "all_users": True, "include_self": True},
        headers=admin_headers,
    )
    assert response.status_code == 200
    body = response.json()
    assert body["created"] == 2
    assert len(body["targets"]) == 2

    member_plans = client.get("/api/plans", headers=other_headers).json()
    assert [plan["title"] for plan in member_plans] == ["全员通知"]

    rollback = client.post(
        "/api/admin/schedule/rollback",
        json={"plan_ids": body["plan_ids"]},
        headers=admin_headers,
    )
    assert rollback.status_code == 200
    assert rollback.json()["deleted"] == 2
    assert client.get("/api/plans", headers=other_headers).json() == []


def test_rollback_ignores_manual_plans(client, admin_headers):
    manual = client.post(
        "/api/plans",
        json={"date": TOMORROW.isoformat(), "title": "手写的计划"},
        headers=admin_headers,
    ).json()
    rollback = client.post(
        "/api/admin/schedule/rollback",
        json={"plan_ids": [manual["id"]]},
        headers=admin_headers,
    )
    assert rollback.json()["deleted"] == 0
    assert len(client.get("/api/plans", headers=admin_headers).json()) == 1
